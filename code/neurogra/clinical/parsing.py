"""Bounded patient-document parsing. Optional libraries report missing capability."""
import csv
import importlib.metadata
from pathlib import Path
from .config import ClinicalConfig
from .intake import file_hash
from .schemas import Attachment, ImageAsset, Issue, ParseResult, SourceRecord
from .utils import dumps, fingerprint, identity


def parse_attachment(attachment: Attachment, config: ClinicalConfig) -> ParseResult:
    path = Path(attachment.resolved_path) if attachment.resolved_path else None
    if attachment.status != "available" or path is None or not path.is_file():
        return ParseResult(attachment_id=attachment.attachment_id, parser="none", status="failed",
                           issues=[Issue(code="attachment_unavailable", stage="parsing", message="附件不可读取")])
    suffix = path.suffix.lower()
    result = ParseResult(attachment_id=attachment.attachment_id, parser=suffix.lstrip("."), status="ok")

    def add(text, locator, context=None, raw=None):
        if len(result.segments) >= config.max_segments:
            raise ValueError("segment_limit")
        result.segments.append(SourceRecord(
            record_id=identity("segment", [attachment.attachment_id, digest, locator]),
            kind="segment", locator={"attachment_id": attachment.attachment_id, **locator},
            text=text, raw_value=raw if raw is not None else text,
            context={**attachment.context, **(context or {})}, path=str(path), content_hash=fingerprint([digest, locator, text])))

    try:
        if path.stat().st_size > config.max_attachment_bytes:
            return ParseResult(attachment_id=attachment.attachment_id, parser="none", status="failed",
                               issues=[Issue(code="file_too_large", stage="parsing", message="附件超过大小上限")])
        digest = file_hash(path)
        if suffix in {".txt", ".md"}:
            result.parser = "utf8"
            for index, line in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), 1):
                if line.strip():
                    add(line, {"line": index})
        elif suffix == ".pdf":
            from pypdf import PdfReader
            result.parser_version = importlib.metadata.version("pypdf")
            reader = PdfReader(path)
            if reader.is_encrypted and not reader.decrypt(""):
                raise ValueError("password_required")
            for page_no, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                if text.strip():
                    add(text, {"page": page_no})
                else:
                    result.issues.append(Issue(code="requires_ocr", stage="parsing",
                                              message=f"第{page_no}页无可提取文本"))
            if result.issues:
                result.status = "partial" if result.segments else "requires_ocr"
        elif suffix == ".docx":
            from docx import Document
            result.parser_version = importlib.metadata.version("python-docx")
            document = Document(path)
            for index, paragraph in enumerate(document.paragraphs):
                if paragraph.text.strip():
                    add(paragraph.text, {"paragraph": index})
            for ti, table in enumerate(document.tables):
                for ri, row in enumerate(table.rows):
                    for ci, cell in enumerate(row.cells):
                        if cell.text.strip():
                            add(cell.text, {"table": ti, "row": ri, "column": ci},
                                {"row_context": [c.text for c in row.cells]})
        elif suffix in {".csv", ".tsv"}:
            with path.open(encoding="utf-8-sig", newline="") as source:
                reader = csv.reader(source, delimiter="\t" if suffix == ".tsv" else ",")
                headers = next(reader, [])
                for ri, row in enumerate(reader, 2):
                    for ci, value in enumerate(row):
                        if value != "":
                            add(value, {"row": ri, "column": ci + 1},
                                {"header": headers[ci] if ci < len(headers) else None, "row_context": row})
        elif suffix == ".xlsx":
            import openpyxl
            result.parser_version = importlib.metadata.version("openpyxl")
            formula_book = openpyxl.load_workbook(path, read_only=True, data_only=False)
            value_book = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in formula_book:
                    headers = [str(c.value) if c.value is not None else "" for c in next(sheet.iter_rows(), [])]
                    values = value_book[sheet.title]
                    # Stream both worksheets in lockstep (no repeated random read in read-only mode).
                    for row, cached_row in zip(sheet.iter_rows(), values.iter_rows()):
                        row_context = [str(c.value) if c.value is not None else "" for c in row]
                        for cell, cached_cell in zip(row, cached_row):
                            if cell.value is None:
                                continue
                            formula = cell.data_type == "f"
                            value = cached_cell.value if formula else cell.value
                            raw = value if isinstance(value, (str, int, float, bool)) or value is None else str(value)
                            context = {"header": headers[cell.column - 1] if cell.column <= len(headers) else "",
                                       "row_context": row_context, "formula": cell.value if formula else None,
                                       "formula_cache_available": value is not None if formula else None}
                            add("" if value is None else str(value),
                                {"sheet": sheet.title, "cell": cell.coordinate}, context, raw)
                            if formula and value is None:
                                result.issues.append(Issue(code="formula_cache_missing", stage="parsing",
                                                          message=f"{sheet.title}!{cell.coordinate}公式无缓存值"))
                if result.issues:
                    result.status = "partial"
            finally:
                formula_book.close()
                value_book.close()
        else:
            result.status = "unsupported"
            result.issues.append(Issue(code="unsupported_format", stage="parsing", message=suffix))
        if result.status == "ok" and not result.segments:
            result.status = "partial"
            result.issues.append(Issue(code="no_content", stage="parsing", message="没有可分析内容"))
    except ImportError as exc:
        result.status = "unsupported"
        result.issues.append(Issue(code="parser_not_installed", stage="parsing", message=str(exc.name)))
    except Exception as exc:
        result.status = "partial" if result.segments else "failed"
        result.issues.append(Issue(code="parse_failed", stage="parsing", message=type(exc).__name__))
    return result


def inspect_image(image: ImageAsset, config: ClinicalConfig) -> ImageAsset:
    updated = image.model_copy(deep=True)
    if not image.path or image.status != "available":
        return updated
    path = Path(image.path)
    try:
        if path.is_dir() or path.suffix.lower() == ".dcm":
            import pydicom
            series = {}
            paths = path.rglob("*") if path.is_dir() else [path]
            scanned = 0
            for file in paths:
                if not file.is_file():
                    continue
                roots = [config.resolve(config.patient_data_root), *[config.resolve(p) for p in config.allowed_roots]]
                if not any(file.resolve().is_relative_to(root) for root in roots):
                    continue
                scanned += 1
                if scanned > config.max_segments:
                    updated.metadata["inspection_truncated"] = True
                    break
                try:
                    ds = pydicom.dcmread(file, stop_before_pixels=True)
                    uid = str(getattr(ds, "SeriesInstanceUID", ""))
                    if not uid:
                        continue
                    row = series.setdefault(uid, {"series_id": uid, "modality": str(getattr(ds, "Modality", "")),
                                                 "date": str(getattr(ds, "StudyDate", "")),
                                                 "description": str(getattr(ds, "SeriesDescription", "")),
                                                 "files": 0})
                    row["files"] += 1
                except (pydicom.errors.InvalidDicomError, OSError):
                    continue
            updated.metadata["series_candidates"] = list(series.values())
            modality = image.modality.upper() if image.modality else None
            aliases = {"MRI": "MR", "PET": "PT"}
            updated.metadata["modality_conflict"] = bool(modality and any(
                row["modality"] != aliases.get(modality, modality) for row in series.values()))
            updated.status = "metadata_available" if series else "unreadable"
        else:
            import SimpleITK as sitk
            reader = sitk.ImageFileReader()
            reader.SetFileName(str(path))
            reader.ReadImageInformation()
            updated.metadata.update({"size": list(reader.GetSize()), "spacing": list(reader.GetSpacing()),
                                     "direction": list(reader.GetDirection())})
            updated.status = "metadata_available"
    except ImportError:
        updated.status = "metadata_reader_unavailable"
    except Exception as exc:
        updated.status = "unreadable"
        updated.metadata["error_type"] = type(exc).__name__
    updated.metadata["classification_executed"] = False
    return updated
