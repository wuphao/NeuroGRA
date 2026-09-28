import asyncio
import json
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from pydantic import ValidationError
from neurogra.clinical.config import ClinicalConfig, BudgetConfig
from neurogra.clinical.schemas import *
from neurogra.clinical.storage import RunStore, BudgetExceeded, VersionConflict
from neurogra.clinical.llm import ModelGateway, ModelFailure
from neurogra.clinical.intake import ingest_patient, resolve_asset
from neurogra.clinical.parsing import parse_attachment
from neurogra.clinical.profiling import profile_case, CandidateBatch, Candidate
from neurogra.clinical.orchestration import MainAgent, Executor, fallback_plan, validate_plan
from neurogra.clinical.retrieval import read_case_record
from neurogra.clinical.utils import parse_time


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.config = ClinicalConfig(project_root=self.root, patient_data_root=Path("."))
        self.store = RunStore(self.root / "clinical.sqlite")
        self.run = self.store.create_run({}, {"llm": 5, "retrieval": 5, "image": 0, "seconds": 30})

    def snapshot(self, raw, gateway=None):
        intake = ingest_patient(raw, self.config)
        return profile_case(intake, [], self.config, gateway)


class IntakeTests(Base):
    def test_arbitrary_fields_roundtrip_and_safe_id(self):
        raw = {"患者ID": "../somewhere", "陌生字段/含~": {"值": [None, 5, "原文"]}}
        result = ingest_patient(raw, self.config)
        self.assertEqual(result.raw_snapshot, raw)
        self.assertTrue(any("~1" in r.locator["json_pointer"] for r in result.records))
        self.assertEqual(result.raw_input_hash, ingest_patient(raw, self.config).raw_input_hash)

    def test_only_id_and_invalid_id(self):
        self.assertFalse(self.snapshot({"患者ID": "x"}).observations)
        for raw in ({}, {"患者ID": 3}, {"患者ID": " "}):
            with self.assertRaises(ValidationError):
                ingest_patient(raw, self.config)

    def test_time_precision_and_invalid_day(self):
        self.assertEqual(parse_time("2026年8月").value, "2026-08")
        self.assertEqual(parse_time("2026-02-30").precision, "text")
        self.assertEqual(parse_time("近两年").precision, "text")

    def test_outside_root_not_read(self):
        path, status = resolve_asset("../secret.txt", self.config)
        self.assertIsNone(path)
        self.assertEqual(status, "forbidden")

    def test_missing_image_keeps_modality_time_path(self):
        intake = ingest_patient({"患者ID": "x", "影像数据": [{"模态": "MRI", "时间": "2026-09", "路径": "missing"}]}, self.config)
        self.assertEqual(intake.images[0].time.precision, "month")
        self.assertEqual(intake.images[0].status, "missing")
        self.assertTrue(intake.images[0].source_refs)

    def test_ordinary_path_text_not_opened(self):
        intake = ingest_patient({"患者ID": "x", "描述": "D:/secret.txt"}, self.config)
        self.assertEqual(intake.attachments, [])

    def test_nan_rejected(self):
        with self.assertRaises(ValueError):
            ingest_patient({"患者ID": "x", "数值": float("nan")}, self.config)


class ParserTests(Base):
    def parse(self, name):
        intake = ingest_patient({"患者ID": "x", "附件": {"路径": name}}, self.config)
        return parse_attachment(intake.attachments[0], self.config)

    def test_text_provenance(self):
        (self.root / "r.txt").write_text("第一行\n第二行", encoding="utf-8")
        parsed = self.parse("r.txt")
        self.assertEqual(len(parsed.segments), 2)
        self.assertEqual(parsed.segments[1].locator["line"], 2)

    def test_csv_provenance(self):
        (self.root / "r.csv").write_text("项目,值,单位\n检查A,12,u\n", encoding="utf-8")
        parsed = self.parse("r.csv")
        self.assertEqual(parsed.segments[1].context["header"], "值")
        self.assertEqual(parsed.segments[1].text, "12")

    def test_blank_pdf_requires_ocr(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(self.root / "blank.pdf")
        self.assertEqual(self.parse("blank.pdf").status, "requires_ocr")

    def test_missing_and_unsupported(self):
        self.assertEqual(self.parse("missing.pdf").status, "failed")
        (self.root / "data.xyz").write_text("sample")
        self.assertEqual(self.parse("data.xyz").status, "unsupported")

    def test_docx_table_location(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("optional docx parser")
        doc = Document()
        doc.add_paragraph("患者叙述")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "20"
        doc.save(self.root / "r.docx")
        parsed = self.parse("r.docx")
        self.assertTrue(any(r.locator.get("table") == 0 for r in parsed.segments))

    def test_xlsx_formula_not_guessed(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("optional xlsx parser")
        book = Workbook()
        book.active.append(["项目", "值"])
        book.active.append(["检查A", "=1+1"])
        book.save(self.root / "r.xlsx")
        parsed = self.parse("r.xlsx")
        self.assertEqual(parsed.status, "partial")
        self.assertTrue(any(i.code == "formula_cache_missing" for i in parsed.issues))


class ProfilingTests(Base):
    def test_missing_not_negative(self):
        snapshot = self.snapshot({"患者ID": "x", "症状": "未提到幻觉", "检验": "未做检查", "病史": "否认幻觉"})
        statuses = {o.quote: o.status for o in snapshot.observations}
        self.assertEqual(statuses["未提到幻觉"], "not_recorded")
        self.assertEqual(statuses["未做检查"], "not_performed")
        self.assertEqual(statuses["否认幻觉"], "absent")

    def test_shared_fact_and_no_lab_routing(self):
        snapshot = self.snapshot({"患者ID": "x", "描述": "吃药需要家属提醒"})
        self.assertEqual(len(snapshot.observations), 1)
        plan = fallback_plan(snapshot)
        self.assertEqual({t.agent_type for t in plan.tasks}, {"history", "cognition"})
        self.assertEqual(validate_plan(plan, snapshot), [])

    def test_unknown_scale_not_named(self):
        snapshot = self.snapshot({"患者ID": "x", "陌生量表": {"总分": 17, "时间": "2026-09"}})
        self.assertFalse(any("MoCA" in o.name for o in snapshot.observations))
        self.assertEqual(snapshot.observations[0].time.precision, "month")

    def test_conflicts_same_day_and_longitudinal(self):
        raw = {"患者ID": "x", "检验": [{"数值": 1, "时间": "2026-09-12"}, {"数值": 2, "时间": "2026-09-12"}]}
        self.assertEqual(len(self.snapshot(raw).conflict_groups), 1)
        raw["检验"][1]["时间"] = "2026-09-13"
        self.assertEqual(self.snapshot(raw).conflict_groups, [])

    def test_unsupported_extraction_rejected(self):
        class Fake:
            def generate_structured(self, *args):
                return CandidateBatch(candidates=[Candidate(record_id="made_up", quote="捏造", domains=["history"], name="症状")])
        snapshot = self.snapshot({"患者ID": "x", "病史": "记忆下降"}, Fake())
        self.assertTrue(any(i.code == "unsupported_extracted_fact" for i in snapshot.issues))
        self.assertFalse(any(o.quote == "捏造" for o in snapshot.observations))

    def test_record_access_version_and_range(self):
        snapshot = self.snapshot({"患者ID": "x", "病史": "记忆下降"})
        record = snapshot.records[-1]
        self.assertEqual(read_case_record(snapshot, record.record_id, 1)["text"], "记忆下降")
        with self.assertRaises(ValueError):
            read_case_record(snapshot, record.record_id, 2)
        with self.assertRaises(ValueError):
            read_case_record(snapshot, record.record_id, 1, -1)

    def test_invalid_task_and_unimplemented_handler(self):
        snapshot = self.snapshot({"患者ID": "x", "病史": "高血压"})
        plan = fallback_plan(snapshot)
        plan.tasks[0].input_refs = ["fabricated"]
        self.assertIn("invalid_input_ref", validate_plan(plan, snapshot))
        results = asyncio.run(Executor(self.store, self.run).execute_tasks(fallback_plan(snapshot), snapshot))
        self.assertEqual(results[0].issues[0].code, "agent_unavailable")

    def test_registered_handler_receives_only_permitted_data(self):
        snapshot = self.snapshot({"患者ID": "x", "病史": "高血压", "检验结果": 3})
        class Handler:
            async def analyze(inner, task, context):
                self.assertFalse(any("laboratory" in o.domains for o in context.snapshot.observations))
                return AgentResult(result_id="r_history", task_id=task.task_id, case_version=1,
                                   status="completed", used_refs=task.input_refs)
        executor = Executor(self.store, self.run)
        executor.register("history", Handler())
        results = asyncio.run(executor.execute_tasks(fallback_plan(snapshot), snapshot))
        self.assertTrue(any(r.status == "completed" for r in results))
        self.assertTrue(any(r.status == "failed" for r in results))


class GatewayStoreTests(Base):
    def test_atomic_budget_persists(self):
        run = self.store.create_run({}, {"llm": 3, "retrieval": 0, "image": 0, "seconds": 30})
        def call(_):
            try:
                return self.store.reserve(run, "llm")[0]
            except BudgetExceeded:
                return None
        with ThreadPoolExecutor(max_workers=8) as pool:
            results = list(pool.map(call, range(10)))
        self.assertEqual(sum(r is not None for r in results), 3)
        reopened = RunStore(self.store.path)
        with self.assertRaises(BudgetExceeded):
            reopened.reserve(run, "llm")
        reopened.mark_interrupted(run)
        self.assertEqual(reopened.summary(run)["calls"][0]["status"], "indeterminate")

    def test_version_compare_and_swap(self):
        self.store.save(self.run, "result", "r", {"value": 1})
        with self.assertRaises(VersionConflict):
            self.store.save(self.run, "result", "r", {"value": 2})
        self.store.save(self.run, "result", "r", {"value": 2}, version=2, expected_version=1)

    def test_format_repair_counted_and_validated(self):
        class Response(Contract):
            ok: bool
        calls = []
        def transport(url, body, timeout):
            calls.append(body)
            return {"message": {"content": "bad" if len(calls) == 1 else '{"ok":true}'}}
        gateway = ModelGateway(self.config.model, self.store, self.run, transport=transport)
        self.assertTrue(gateway.generate_structured("test", {}, Response).ok)
        self.assertEqual(len(calls), 2)
        self.assertEqual(sum(r["n"] for r in self.store.summary(self.run)["calls"]), 2)

    def test_input_not_silently_truncated(self):
        config = self.config.model.model_copy(update={"max_input_chars": 2})
        with self.assertRaises(ModelFailure):
            ModelGateway(config, self.store, self.run).generate_structured("test", {"large": "payload"}, Contract)

    def test_main_planner_fallback(self):
        class Failing:
            def generate_structured(self, *args):
                raise ModelFailure("unavailable")
        plan = MainAgent(Failing()).plan(self.snapshot({"患者ID": "x", "病史": "高血压"}))
        self.assertEqual(plan.mode, "deterministic_fallback")
        self.assertTrue(plan.issues)


if __name__ == "__main__":
    unittest.main()
