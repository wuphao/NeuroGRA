"""Build the standalone NeuroGRA algorithm design with native Word mathematics.

This is document authoring, not an implementation of the retrieval algorithm.
Uses the selected desktop bundled Python and the installed Office MathML XSL.
"""
from __future__ import annotations

import importlib.util
from copy import deepcopy
import json
import re
import sys
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'docs/知识图谱设计/NeuroGRA_多模态证据检索与缺口补偿算法设计_v1.0.md'
ASSETS = ROOT / 'docs/assets/neurogra_gap_method_v1'
OUT = ROOT / 'output/documents/NeuroGRA_多模态证据检索与缺口补偿算法设计_v1.0.docx'
BUILD = ROOT / 'output/.build/neurogra_method_document_v1'
MATH_XSL = Path('C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL')
MML = 'http://www.w3.org/1998/Math/MathML'
TRANSFORM = etree.XSLT(etree.parse(str(MATH_XSL)))


def node(tag, *children, text=None, **attrs):
    e = etree.Element(f'{{{MML}}}{tag}', **attrs)
    if text is not None:
        e.text = text
    for ch in children:
        e.append(ch)
    return e


COMMANDS = {
    'mathcal': 'style', 'mathrm': 'style', 'operatorname': 'roman', 'text': 'roman',
    'mathbb': 'double', 'frac': 'frac', 'sum': 'sum', 'bigcup': 'bigcup',
    'overline': 'bar', 'neg': '¬', 'land': '∧', 'lor': '∨', 'mid': '|',
    'cdot': '·', 'lambda': 'λ', 'phi': 'φ', 'psi': 'ψ', 'sigma': 'σ',
    'kappa': 'κ', 'epsilon': 'ε', 'pi': 'π', 'mu': 'μ', 'eta': 'η',
    'delta': 'δ', 'tau': 'τ', 'alpha': 'α', 'Lambda': 'Λ', 'cup': '∪',
    'in': '∈', 'subseteq': '⊆', 'le': '≤', 'varnothing': '∅', 'ne': '≠',
    'exp': 'exp', 'log': 'log', 'left': 'skip', 'right': 'skip', 'ell': 'ℓ', 'qquad': 'space', 'exists': '∃',
}


class LatexMath:
    """Small explicit parser for the notation actually used by this document."""
    def __init__(self, text):
        self.text = text
        self.pos = 0

    def ws(self):
        while self.pos < len(self.text) and self.text[self.pos].isspace():
            self.pos += 1

    def group_raw(self):
        self.ws()
        if self.text[self.pos] != '{':
            raw = self.text[self.pos]
            self.pos += 1
            return raw
        start = self.pos + 1
        self.pos += 1
        level = 1
        while self.pos < len(self.text):
            c = self.text[self.pos]
            self.pos += 1
            if c == '{': level += 1
            if c == '}': level -= 1
            if level == 0: return self.text[start:self.pos - 1]
        raise ValueError('Unclosed raw group')

    def atom(self):
        self.ws()
        if self.pos >= len(self.text):
            raise ValueError('Expected atom')
        c = self.text[self.pos]
        self.pos += 1
        if c == '{':
            e = self.sequence('}')
        elif c == '\\':
            if self.pos < len(self.text) and self.text[self.pos] in '{}|,;!':
                escaped = self.text[self.pos]
                self.pos += 1
                return node('mo', text=('‖' if escaped == '|' else escaped) if escaped not in ',;!' else '')
            if self.pos < len(self.text) and self.text[self.pos] == ' ':
                self.pos += 1
                return node('mspace', width='.2em')
            m = re.match(r'[A-Za-z]+', self.text[self.pos:])
            if not m: raise ValueError(f'Invalid command at {self.pos}')
            cmd = m.group()
            self.pos += len(cmd)
            if cmd not in COMMANDS: raise ValueError(f'Unsupported command {cmd}')
            action = COMMANDS[cmd]
            if action == 'skip': return self.atom()
            if action == 'space': return node('mspace', width='1em')
            if action in ('style', 'roman', 'double'):
                raw = self.group_raw()
                if action == 'roman': e = node('mi', text=raw, mathvariant='normal')
                elif action == 'double': e = node('mi', text=raw, mathvariant='double-struck')
                elif cmd == 'mathcal': e = node('mi', text=raw, mathvariant='script')
                else: e = node('mi', text=raw, mathvariant='normal')
            elif action == 'frac':
                e = node('mfrac', self.atom(), self.atom())
            elif action == 'bar':
                e = node('mover', self.atom(), node('mo', text='¯'), accent='true')
            elif action in ('sum', 'bigcup'):
                e = node('mo', text='∑' if action == 'sum' else '⋃', largeop='true')
            else:
                kind = 'mi' if action in 'φψσκεπμηδταλΛ' or action in ('exp', 'log') else 'mo'
                e = node(kind, text=action, **({'mathvariant': 'normal'} if action in ('exp', 'log') else {}))
        elif c.isdigit():
            tail = re.match(r'\d*', self.text[self.pos:]).group()
            self.pos += len(tail)
            e = node('mn', text=c + tail)
        else:
            e = node('mi' if c.isalpha() else 'mo', text=c)
        return e

    def scripted(self):
        e = self.atom()
        sub, sup = None, None
        while True:
            self.ws()
            if self.pos >= len(self.text) or self.text[self.pos] not in '_^': break
            which = self.text[self.pos]
            self.pos += 1
            val = self.atom()
            if which == '_': sub = val
            else: sup = val
        if sub is not None and sup is not None: return node('msubsup', e, sub, sup)
        if sub is not None: return node('msub', e, sub)
        if sup is not None: return node('msup', e, sup)
        return e

    def sequence(self, end=None):
        parts = []
        while True:
            self.ws()
            if self.pos >= len(self.text):
                if end: raise ValueError(f'Missing {end}')
                break
            if end and self.text[self.pos] == end:
                self.pos += 1
                break
            parts.append(self.scripted())
        return node('mrow', *parts)


def omml(latex):
    latex = re.sub(r'\\tag\{\d+\}', '', latex)
    math = node('math', LatexMath(latex).sequence(), display='inline')
    converted = TRANSFORM(math).getroot()
    if converted.tag.endswith('oMathPara'):
        converted = converted.find(qn('m:oMath'))
    if converted is None: raise ValueError(f'Cannot convert {latex}')
    # Office's MathML XSL creates a separate n-ary operator with an empty operand.
    # Attach the following term to its m:e so renderers do not show a placeholder.
    ns = {'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    for op in reversed(converted.xpath('.//m:nary', namespaces=ns)):
        operand = op.find(qn('m:e'))
        if operand is None or len(operand):
            continue
        parent = op.getparent()
        depth = 0
        while op.getnext() is not None:
            following = op.getnext()
            if following.tag != qn('m:r'):
                operand.append(following)
                continue
            tnode = following.find(qn('m:t'))
            txt = '' if tnode is None else tnode.text or ''
            stop = None
            for j, ch in enumerate(txt):
                if depth == 0 and ch in '+-=,.;≤≥)}]':
                    stop = j
                    break
                if ch in '({[': depth += 1
                elif ch in ')}]': depth -= 1
            if stop is None:
                operand.append(following)
            elif stop == 0:
                break
            else:
                prefix = deepcopy(following)
                prefix.find(qn('m:t')).text = txt[:stop]
                tnode.text = txt[stop:]
                operand.append(prefix)
                break
        if not len(operand):
            raise ValueError(f'Unbound n-ary operand: {latex}')
    return converted


def font_run(run, size=11.5, bold=False, east='宋体', latin='Times New Roman'):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'), east)


def inline(par, text, size=11.5):
    parts = re.split(r'(\\\(.*?\\\))', text)
    for part in parts:
        if part.startswith('\\(') and part.endswith('\\)'):
            par._p.append(omml(part[2:-2]))
        else:
            for j, s in enumerate(re.split(r'(\*\*.*?\*\*)', part)):
                font_run(par.add_run(s[2:-2] if s.startswith('**') else s), size,
                         bold=s.startswith('**'))


def generate_figures():
    spec = importlib.util.spec_from_file_location('diagram_base', ROOT / 'code/tools/render_graph_method_figures.py')
    base = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(base)
    base.OUT = ASSETS
    C, blue, green, orange = base.Canvas, base.BLUE, base.GREEN, base.ORANGE

    c = C(1200, 700)
    c.text(600, 35, 'NeuroGRA 多模态证据检索与缺口补偿', 34, bold=True)
    c.text(35, 77, '离线构建', 27, blue, True, 'lm')
    c.box(30, 103, 350, 105, '固定版本医学文献', ['正文 · 图像 · 表格 · 公式'], title_size=30, body_size=27)
    c.box(425, 103, 350, 105, '抽取与领域知识汇整', ['实体对齐 · 条件保留 · 来源'], title_size=30, body_size=27)
    c.box(820, 103, 350, 105, '图谱与检索索引', ['语义索引 · 类型统计 · 评分'], title_size=30, body_size=27)
    c.arrow([(380,155),(425,155)])
    c.arrow([(775,155),(820,155)])
    c.text(600, 245, '离线产物供在线检索与候选探索使用', 28, base.MUTED)
    c.text(35, 285, '在线查询', 27, blue, True, 'lm')
    c.box(30, 315, 350, 110, '需求规划与初始检索', ['明确字段 · 图与语义候选'], title_size=30, body_size=27)
    c.box(425, 315, 350, 110, '实际证据覆盖检查', ['核对字段 · 范围 · 条件依赖'], color=green, title_size=30, body_size=27)
    c.box(820, 315, 350, 110, '缺口分类与候选探索', ['合法关系 · 已知或新实体'], color=orange, fill='#FFF7ED', title_size=30, body_size=27)
    c.arrow([(380,370),(425,370)])
    c.arrow([(775,370),(820,370)], orange)
    c.box(820, 520, 350, 110, '来源回查与剩余补偿', ['本地原文 · 可选外部文献'], color=orange, fill='#FFF7ED', title_size=30, body_size=27)
    c.box(425, 520, 350, 110, '核验与证据组装', ['定位 · 语义 · 范围 · 上下文'], color=green, title_size=30, body_size=27)
    c.box(30, 520, 350, 110, '带引用回答与缺口', ['保留条件 · 不确定性 · 来源'], color=green, title_size=30, body_size=27)
    c.arrow([(995,425),(995,520)], orange)
    c.arrow([(820,575),(775,575)], green)
    c.arrow([(600,520),(600,425)], green)
    c.text(690,474,'重新检查',23,green)
    c.arrow([(460,425),(460,474),(205,474),(205,520)],green)
    c.text(235,456,'已充分或停止',24,green)
    c.text(600,674,'候选只提供搜索方向  事实回答必须有可回看的原始来源',28,base.MUTED)
    c.save('overview')

    c = C(1200, 550)
    c.text(600,38,'候选与证据的状态转换',34,bold=True)
    c.box(30,115,350,120,'运行时探索候选',['关系与实体候选ID','保留目标需求和范围'],color=orange,fill='#FFF7ED',title_size=30,body_size=27)
    c.box(425,115,350,120,'原始来源核验',['定位与关系语义','范围与解释上下文'],title_size=30,body_size=27)
    c.box(820,115,350,120,'来源支持的证据池',['可用于本次引用回答','保留来源与审核状态'],color=green,title_size=30,body_size=27)
    c.arrow([(380,175),(425,175)],orange)
    c.arrow([(775,175),(820,175)],green)
    c.box(425,335,350,120,'反驳 冲突或未知',['保留证据或失败原因','不充当肯定支持'],color=base.GRAY,fill='#F5F6F8',title_size=30,body_size=27)
    c.box(820,335,350,120,'审核后更新正式图谱',['正式发布独立执行','随后更新索引与向量'],color=green,title_size=30,body_size=27)
    c.arrow([(600,235),(600,335)],base.GRAY)
    c.arrow([(995,235),(995,335)],green,dashed=True)
    c.text(600,510,'预测分数不能代替原文核验  查询回答不自动发布新知识',28,base.MUTED)
    c.save('evidence_states')

    c = C(1200, 680)
    c.text(600,40,'证据缺口与补偿动作',34,bold=True)
    pairs=[('连接疑似缺失','合法候选与原文回查'),('实体尚未入图','来源发现与身份对齐'),('必要条件不完整','依赖展开与上下文补查'),('本地来源仍不足','受控外部检索或保留缺口')]
    for i,(l,r) in enumerate(pairs):
        y=100+i*125
        c.box(30,y,430,95,l,[],color=orange,fill='#FFF7ED',title_size=31)
        c.box(610,y,560,95,r,[],color=blue,title_size=31)
        c.arrow([(460,y+47),(610,y+47)],orange)
    c.text(600,620,'来源冲突分离核对  上下文容量不足优先重组证据',28,base.MUTED)
    c.text(600,657,'所有新增内容先核验  剩余未知不转为否定事实',28,base.MUTED)
    c.save('gap_routes')


def set_style(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.78)
    sec.left_margin = sec.right_margin = Inches(.8)
    sec.header_distance = sec.footer_distance = Inches(.34)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Times New Roman', Pt(11.5)
    normal.font.color.rgb = RGBColor(0,0,0)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.widow_control = True
    for name,size in [('Title',21),('Subtitle',12),('Heading 1',15),('Heading 2',12.5),('Heading 3',12)]:
        sty=doc.styles[name]
        sty.font.name='Times New Roman'
        sty.font.size=Pt(size)
        sty.font.color.rgb=RGBColor(0,0,0)
        sty._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'黑体' if name!='Subtitle' else '宋体')
        sty.paragraph_format.space_before=Pt(13 if name.startswith('Heading') else 5)
        sty.paragraph_format.space_after=Pt(7)
        if name.startswith('Heading'): sty.paragraph_format.keep_with_next=True
        borders=sty._element.find('.//' + qn('w:pBdr'))
        if borders is not None: borders.getparent().remove(borders)
    header=sec.header.paragraphs[0]
    font_run(header.add_run('NeuroGRA 算法设计'),9,east='宋体')
    header.paragraph_format.space_after=Pt(0)
    footer=sec.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    font_run(footer.add_run('第 '),9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    font_run(footer.add_run(' 页'),9)
    settings=doc.settings.element
    mathpr=settings.find(qn('m:mathPr'))
    if mathpr is None: mathpr=OxmlElement('m:mathPr');settings.append(mathpr)
    mathfont=OxmlElement('m:mathFont');mathfont.set(qn('m:val'),'Cambria Math');mathpr.append(mathfont)


def add_table(doc, rows):
    rows=[[c.strip() for c in row.strip('|').split('|')] for row in rows]
    rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',c) for c in r)]
    table=doc.add_table(rows=1,cols=len(rows[0]))
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    widths={2:[1.48,5.42],3:[1.3,2.9,2.7]}.get(len(rows[0]),[6.9/len(rows[0])]*len(rows[0]))
    for col,w in zip(table.columns,widths): col.width=Inches(w)
    for idx,row in enumerate(rows):
        cells=table.rows[0].cells if idx==0 else table.add_row().cells
        for ci,(cell,text) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[ci])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.12
            inline(p,text,10.5)
            tcpr=cell._tc.get_or_add_tcPr()
            shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E8EDF2' if idx==0 else ('F7F9FB' if idx%2==0 else 'FFFFFF'));tcpr.append(shading)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                m=OxmlElement('w:'+side);m.set(qn('w:w'),'85');m.set(qn('w:type'),'dxa');margins.append(m)
            tcpr.append(margins)
            if idx==0:
                for run in p.runs: run.bold=True
        if idx==0:
            repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
        prevent=OxmlElement('w:cantSplit');table.rows[idx]._tr.get_or_add_trPr().append(prevent)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)


def build_doc():
    text=SOURCE.read_text(encoding='utf-8')
    doc=Document();set_style(doc)
    lines=text.splitlines();i=0;eq_count=0;image_count=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line: continue
        if line=='\\[':
            parts=[]
            while lines[i].strip()!='\\]': parts.append(lines[i].strip());i+=1
            i+=1;latex=' '.join(parts);number=re.search(r'\\tag\{(\d+)\}',latex).group(1)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(7);p.paragraph_format.space_after=Pt(8)
            p.paragraph_format.keep_with_next=True
            p._p.append(omml(latex));font_run(p.add_run('   ('+number+')'),10.5)
            eq_count+=1
        elif line.startswith('```'):
            code_index = 0
            while not lines[i].strip().startswith('```'):
                code=lines[i];i+=1
                p=doc.add_paragraph();p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(2)
                p.paragraph_format.keep_with_next = code_index < 6
                font_run(p.add_run(code),10.5,east='宋体',latin='Consolas')
                code_index += 1
            i+=1
        elif line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
            add_table(doc,rows)
        elif line.startswith('!['):
            path=re.search(r'\]\((.+)\)$',line).group(1)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str((SOURCE.parent/path).resolve()),width=Inches(6.8))
            image_count+=1
        elif line.startswith('# '):
            p=doc.add_paragraph(line[2:],style='Title');p.paragraph_format.keep_with_next=True
        elif line.startswith('## '):
            doc.add_paragraph(line[3:],style='Heading 1')
        elif line.startswith('### '):
            p = doc.add_paragraph(line[4:],style='Heading 2')
            if line.startswith('### 12.2 '):
                p.paragraph_format.page_break_before = True
        else:
            p=doc.add_paragraph()
            if line.startswith('图') and re.match(r'图\d',line):
                inline(p,line,10);p.paragraph_format.space_after=Pt(9)
            elif line.startswith('面向不完备'):
                p.style='Subtitle';inline(p,line,12)
            elif line.startswith('版本 v'):
                inline(p,line,10);p.paragraph_format.space_after=Pt(12)
            else:
                inline(p,line)
                # Keep an equation's introduction with the numbered equation.
                next_nonempty = next((x.strip() for x in lines[i:] if x.strip()), '')
                if next_nonempty == '\\[':
                    p.paragraph_format.keep_with_next = True
    OUT.parent.mkdir(parents=True,exist_ok=True);BUILD.mkdir(parents=True,exist_ok=True)
    doc.core_properties.title='NeuroGRA 多模态证据检索与缺口补偿算法设计'
    doc.core_properties.subject='面向不完备医学知识图谱的需求驱动探索与来源约束'
    doc.core_properties.author='NeuroGRA'
    doc.save(OUT)
    with ZipFile(OUT) as z:
        xml=etree.fromstring(z.read('word/document.xml'))
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
        alltext=''.join(xml.xpath('//w:t/text()',namespaces=ns))
        assert not re.search(r'RAG[- ]?Anything|LightRAG|Explore-on-Graph|\bXoG\b',alltext,re.I)
        assert '\\[' not in alltext and '\\(' not in alltext
        math_count=len(xml.xpath('//m:oMath',namespaces=ns))
        assert eq_count==18 and image_count==3 and math_count>=18
        assert not xml.xpath('//m:nary/m:e[not(*)]', namespaces=ns)
    record={'output':str(OUT),'display_equations':eq_count,'native_math_objects':math_count,'figures':image_count,
            'tables':len(doc.tables),'main_body_external_framework_names':0,'status':'Design only; no experimental results'}
    (BUILD/'build_validation.json').write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(record,ensure_ascii=False))


if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    generate_figures()
    build_doc()
